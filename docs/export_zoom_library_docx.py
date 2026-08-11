from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


COLLECTIONS = (
    ("big_book", "Большая книга", "Большая книга.docx", "excerpt"),
    ("twelve_twelve", "12 шагов и 12 традиций", "12 шагов и 12 традиций.docx", "excerpt"),
    ("living_sober", "Жить трезвыми", "Жить трезвыми.docx", "excerpt"),
    ("daily_reflections", "Ежедневные размышления", "Ежедневные размышления.docx", "date"),
    ("as_bill_sees_it", "Как это видит Билл", "Как это видит Билл.docx", "excerpt"),
    ("game_questions", "Вопросы игры", "500 вопросов.docx", "question"),
)

BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
GRAY = RGBColor(0x66, 0x66, 0x66)
BLACK = RGBColor(0x00, 0x00, 0x00)


def set_font(run, name: str, size: float | None = None, *, bold: bool | None = None,
             color: RGBColor | None = None, italic: bool | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def set_style_font(style, name: str, size: float, color: RGBColor, *, bold: bool = False) -> None:
    style.font.name = name
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    style.font.size = Pt(size)
    style.font.color.rgb = color
    style.font.bold = bold


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    set_font(run, "Calibri", 9, color=GRAY)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, separate, text, end):
        run._r.append(element)


def configure_document(doc: Document, title: str) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    set_style_font(normal, "Calibri", 11, BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.widow_control = True

    h1 = doc.styles["Heading 1"]
    set_style_font(h1, "Calibri", 16, BLUE, bold=True)
    h1.paragraph_format.space_before = Pt(18)
    h1.paragraph_format.space_after = Pt(10)
    h1.paragraph_format.keep_with_next = True

    h2 = doc.styles["Heading 2"]
    set_style_font(h2, "Calibri", 13, BLUE, bold=True)
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(7)
    h2.paragraph_format.keep_with_next = True

    h3 = doc.styles["Heading 3"]
    set_style_font(h3, "Calibri", 12, DARK_BLUE, bold=True)
    h3.paragraph_format.space_before = Pt(10)
    h3.paragraph_format.space_after = Pt(5)
    h3.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    set_font(header.add_run(title), "Calibri", 9, color=GRAY)
    add_page_field(section.footer.paragraphs[0])


def clean_text(value: object) -> str:
    text = str(value or "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"\A---\n.*?\n---\n?", "", text, flags=re.DOTALL)
    text = re.sub(r"\n?\s*<!--\s*zoom-part\s*-->\s*\n?", "\n\n", text, flags=re.IGNORECASE)
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def sort_key(key: str, kind: str):
    if kind == "date":
        day, month = (int(part) for part in key.split(".", 1))
        return month, day
    try:
        return int(key)
    except ValueError:
        return key


def entry_heading(key: str, kind: str) -> str:
    if kind == "date":
        return key
    if kind == "question":
        return f"Вопрос №{int(key)}"
    return f"Отрывок №{int(key)}"


def add_entry(doc: Document, key: str, text: str, kind: str) -> None:
    doc.add_paragraph(entry_heading(key, kind), style="Heading 2")
    blocks = [part.strip() for part in re.split(r"\n\s*\n", clean_text(text)) if part.strip()]
    if not blocks:
        blocks = [""]
    for index, block in enumerate(blocks):
        paragraph = doc.add_paragraph(style="Normal")
        paragraph.paragraph_format.keep_together = False
        paragraph.paragraph_format.space_after = Pt(12 if index == len(blocks) - 1 else 6)
        lines = block.split("\n")
        for line_index, line in enumerate(lines):
            if line_index:
                paragraph.add_run().add_break()
            paragraph.add_run(line)


def build_document(title: str, entries: dict[str, object], kind: str, output_path: Path) -> None:
    doc = Document()
    configure_document(doc, title)

    title_paragraph = doc.add_paragraph()
    title_paragraph.paragraph_format.space_before = Pt(0)
    title_paragraph.paragraph_format.space_after = Pt(6)
    title_run = title_paragraph.add_run(title)
    set_font(title_run, "Calibri", 22, bold=True, color=BLACK)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(16)
    set_font(subtitle.add_run(f"Нарезанная база: {len(entries)} материалов"), "Calibri", 10, color=GRAY)

    for key in sorted(entries, key=lambda value: sort_key(value, kind)):
        value = entries[key]
        text = value.get("text", "") if isinstance(value, dict) else value
        add_entry(doc, str(key), str(text), kind)

    core = doc.core_properties
    core.title = title
    core.subject = "Нарезанная база Нафани для Zoom"
    core.author = "Группа АА «Почти нормальные»"
    core.keywords = "АА, Нафаня, Zoom, отрывки"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("library_json", type=Path)
    parser.add_argument("output_dir", type=Path)
    args = parser.parse_args()

    payload = json.loads(args.library_json.read_text(encoding="utf-8"))
    collections = payload.get("collections", {})
    manifest = []
    for collection_id, fallback_title, filename, kind in COLLECTIONS:
        collection = collections.get(collection_id) or {}
        entries = collection.get("entries") or {}
        if not isinstance(entries, dict) or not entries:
            raise SystemExit(f"Collection {collection_id} is empty or missing")
        title = str(collection.get("label") or fallback_title)
        output_path = args.output_dir / filename
        build_document(title, entries, kind, output_path)
        manifest.append({"id": collection_id, "title": title, "count": len(entries), "file": filename})

    (args.output_dir / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )


if __name__ == "__main__":
    main()
